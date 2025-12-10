"""
Servicio para generar reportes exportables en PDF y Excel.
"""
import io
from datetime import datetime, timedelta
from django.http import HttpResponse
from django.utils import timezone
from django.db.models import Count, Avg, Q, F, Sum
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import Class, UserClassReservation, ClassAttendance, ClassWaitlist
from .services import ClassDashboardService
from django.contrib.auth import get_user_model

User = get_user_model()


class ClassExportService:
    """
    Servicio para exportar reportes de clases en PDF y Excel.
    """
    
    @staticmethod
    def export_attendance_report(class_id, format='pdf'):
        """
        Exporta reporte de asistencia de una clase específica.
        
        Args:
            class_id: ID de la clase
            format: 'pdf' o 'excel'
        """
        try:
            class_obj = Class.objects.get(id=class_id)
        except Class.DoesNotExist:
            raise ValueError(f"Clase con ID {class_id} no encontrada")
        
        reservations = UserClassReservation.objects.filter(
            class_reserved=class_obj,
            is_cancelled=False
        ).select_related('user')
        
        attendance_records = ClassAttendance.objects.filter(
            class_reserved=class_obj
        ).select_related('user')
        
        # Crear diccionario de asistencia
        attendance_dict = {record.user_id: record.attended for record in attendance_records}
        
        if format == 'pdf':
            return ClassExportService._generate_attendance_pdf(
                class_obj, reservations, attendance_dict
            )
        else:
            return ClassExportService._generate_attendance_excel(
                class_obj, reservations, attendance_dict
            )
    
    @staticmethod
    def export_instructor_statistics(instructor_id, period_months=6, format='pdf'):
        """
        Exporta estadísticas de un instructor.
        
        Args:
            instructor_id: ID del instructor
            period_months: Período en meses
            format: 'pdf' o 'excel'
        """
        try:
            instructor = User.objects.get(id=instructor_id)
        except User.DoesNotExist:
            raise ValueError(f"Instructor con ID {instructor_id} no encontrado")
        
        today = timezone.now().date()
        period_start = today.replace(day=1) - timedelta(days=(period_months - 1) * 30)
        
        classes = Class.objects.filter(
            instructor_id=instructor_id,
            date__date__gte=period_start,
            date__date__lte=today
        )
        
        # Calcular estadísticas
        total_classes = classes.count()
        total_reservations = UserClassReservation.objects.filter(
            class_reserved__in=classes,
            is_cancelled=False
        ).count()
        
        total_attendance = ClassAttendance.objects.filter(
            class_reserved__in=classes,
            attended=True
        ).count()
        
        cancelled_classes = classes.filter(is_cancelled=True).count()
        
        # Ocupación promedio
        occupancy_data = []
        for cls in classes:
            if cls.max_students > 0:
                occupancy = (cls.reservation_count / cls.max_students) * 100
                occupancy_data.append({
                    'class': cls,
                    'occupancy': occupancy
                })
        
        avg_occupancy = sum(d['occupancy'] for d in occupancy_data) / len(occupancy_data) if occupancy_data else 0
        
        # Estadísticas por tipo de clase
        class_type_stats = classes.values('class_type').annotate(
            count=Count('id'),
            total_reservations=Count('userclassreservation', filter=Q(userclassreservation__is_cancelled=False))
        )
        
        if format == 'pdf':
            return ClassExportService._generate_instructor_stats_pdf(
                instructor, period_months, total_classes, total_reservations,
                total_attendance, cancelled_classes, avg_occupancy, class_type_stats, classes
            )
        else:
            return ClassExportService._generate_instructor_stats_excel(
                instructor, period_months, total_classes, total_reservations,
                total_attendance, cancelled_classes, avg_occupancy, class_type_stats, classes
            )
    
    @staticmethod
    def export_occupancy_analysis(start_date, end_date, format='pdf'):
        """
        Exporta análisis de ocupación de clases en un período.
        
        Args:
            start_date: Fecha de inicio
            end_date: Fecha de fin
            format: 'pdf' o 'excel'
        """
        classes = Class.objects.filter(
            date__date__gte=start_date,
            date__date__lte=end_date
        ).select_related('instructor').order_by('date')
        
        occupancy_data = []
        for cls in classes:
            if cls.max_students > 0:
                occupancy_rate = (cls.reservation_count / cls.max_students) * 100
                occupancy_data.append({
                    'class': cls,
                    'occupancy_rate': occupancy_rate,
                    'available_spots': cls.max_students - cls.reservation_count
                })
        
        # Estadísticas agregadas
        total_classes = classes.count()
        total_capacity = sum(cls.max_students for cls in classes)
        total_reservations = sum(cls.reservation_count for cls in classes)
        avg_occupancy = sum(d['occupancy_rate'] for d in occupancy_data) / len(occupancy_data) if occupancy_data else 0
        
        # Por tipo de clase
        type_stats = {}
        for cls in classes:
            cls_type = cls.get_class_type_display()
            if cls_type not in type_stats:
                type_stats[cls_type] = {'total': 0, 'reservations': 0, 'capacity': 0}
            type_stats[cls_type]['total'] += 1
            type_stats[cls_type]['reservations'] += cls.reservation_count
            type_stats[cls_type]['capacity'] += cls.max_students
        
        if format == 'pdf':
            return ClassExportService._generate_occupancy_pdf(
                start_date, end_date, occupancy_data, total_classes,
                total_capacity, total_reservations, avg_occupancy, type_stats
            )
        else:
            return ClassExportService._generate_occupancy_excel(
                start_date, end_date, occupancy_data, total_classes,
                total_capacity, total_reservations, avg_occupancy, type_stats
            )
    
    @staticmethod
    def export_waitlist_report(start_date=None, end_date=None, format='pdf'):
        """
        Exporta reporte de lista de espera.
        
        Args:
            start_date: Fecha de inicio (opcional)
            end_date: Fecha de fin (opcional)
            format: 'pdf' o 'excel'
        """
        waitlist_entries = ClassWaitlist.objects.filter(
            status='waiting'
        ).select_related('user', 'class_reserved')
        
        if start_date:
            waitlist_entries = waitlist_entries.filter(
                class_reserved__date__date__gte=start_date
            )
        if end_date:
            waitlist_entries = waitlist_entries.filter(
                class_reserved__date__date__lte=end_date
            )
        
        waitlist_entries = waitlist_entries.order_by('class_reserved__date', 'position')
        
        # Estadísticas
        total_entries = waitlist_entries.count()
        unique_users = waitlist_entries.values('user').distinct().count()
        unique_classes = waitlist_entries.values('class_reserved').distinct().count()
        
        # Por clase
        class_stats = {}
        for entry in waitlist_entries:
            cls_id = entry.class_reserved_id
            if cls_id not in class_stats:
                class_stats[cls_id] = {
                    'class': entry.class_reserved,
                    'count': 0,
                    'entries': []
                }
            class_stats[cls_id]['count'] += 1
            class_stats[cls_id]['entries'].append(entry)
        
        if format == 'pdf':
            return ClassExportService._generate_waitlist_pdf(
                waitlist_entries, total_entries, unique_users, unique_classes, class_stats
            )
        else:
            return ClassExportService._generate_waitlist_excel(
                waitlist_entries, total_entries, unique_users, unique_classes, class_stats
            )
    
    @staticmethod
    def export_cancellation_statistics(period_months=6, instructor_id=None, start_date=None, end_date=None, format='pdf'):
        """
        Exporta estadísticas de cancelaciones.
        
        Args:
            period_months: Período en meses
            instructor_id: ID del instructor (opcional)
            start_date: Fecha de inicio (opcional)
            end_date: Fecha de fin (opcional)
            format: 'pdf' o 'excel'
        """
        analysis = ClassDashboardService.get_cancellation_analysis(
            period_months=period_months,
            instructor_id=instructor_id,
            start_date=start_date,
            end_date=end_date
        )
        
        if format == 'pdf':
            return ClassExportService._generate_cancellation_pdf(analysis)
        else:
            return ClassExportService._generate_cancellation_excel(analysis)
    
    # ========== Métodos privados para generar PDF ==========
    
    @staticmethod
    def _generate_attendance_pdf(class_obj, reservations, attendance_dict):
        """Genera PDF de reporte de asistencia."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        story = []
        styles = getSampleStyleSheet()
        
        # Título
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#1e40af'),
            spaceAfter=30,
            alignment=1  # Centrado
        )
        story.append(Paragraph("Reporte de Asistencia", title_style))
        story.append(Spacer(1, 0.2*inch))
        
        # Información de la clase
        info_data = [
            ['Clase:', class_obj.name],
            ['Fecha:', class_obj.date.strftime('%d/%m/%Y %H:%M')],
            ['Instructor:', class_obj.instructor.get_full_name() if class_obj.instructor else 'N/A'],
            ['Capacidad:', f"{class_obj.reservation_count}/{class_obj.max_students}"],
        ]
        info_table = Table(info_data, colWidths=[2*inch, 4*inch])
        info_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f3f4f6')),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
        ]))
        story.append(info_table)
        story.append(Spacer(1, 0.3*inch))
        
        # Tabla de asistencia
        attendance_data = [['Estudiante', 'Email', 'Asistió', 'Fecha Registro']]
        for reservation in reservations:
            attended = attendance_dict.get(reservation.user_id, None)
            attended_text = 'Sí' if attended else 'No' if attended is False else 'No registrado'
            attendance_data.append([
                reservation.user.get_full_name() or reservation.user.email,
                reservation.user.email,
                attended_text,
                reservation.created_at.strftime('%d/%m/%Y %H:%M') if reservation.created_at else 'N/A'
            ])
        
        attendance_table = Table(attendance_data, colWidths=[2*inch, 2.5*inch, 1*inch, 1.5*inch])
        attendance_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 11),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')]),
        ]))
        story.append(attendance_table)
        
        # Resumen
        story.append(Spacer(1, 0.3*inch))
        total = len(reservations)
        attended_count = sum(1 for v in attendance_dict.values() if v)
        no_show_count = sum(1 for v in attendance_dict.values() if v is False)
        not_recorded = total - len(attendance_dict)
        
        summary_data = [
            ['Total Reservas:', str(total)],
            ['Asistieron:', str(attended_count)],
            ['No Asistieron:', str(no_show_count)],
            ['No Registrado:', str(not_recorded)],
        ]
        summary_table = Table(summary_data, colWidths=[2*inch, 1*inch])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f3f4f6')),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ]))
        story.append(summary_table)
        
        # Pie de página
        story.append(Spacer(1, 0.3*inch))
        footer = Paragraph(
            f"Generado el: {timezone.now().strftime('%d/%m/%Y %H:%M')}",
            styles['Normal']
        )
        story.append(footer)
        
        doc.build(story)
        buffer.seek(0)
        
        response = HttpResponse(buffer, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="reporte_asistencia_{class_obj.id}_{datetime.now().strftime("%Y%m%d")}.pdf"'
        return response
    
    @staticmethod
    def _generate_instructor_stats_pdf(instructor, period_months, total_classes, total_reservations,
                                       total_attendance, cancelled_classes, avg_occupancy,
                                       class_type_stats, classes):
        """Genera PDF de estadísticas de instructor."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        story = []
        styles = getSampleStyleSheet()
        
        # Título
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#1e40af'),
            spaceAfter=30,
            alignment=1
        )
        story.append(Paragraph("Estadísticas de Instructor", title_style))
        story.append(Spacer(1, 0.2*inch))
        
        # Información del instructor
        info_data = [
            ['Instructor:', instructor.get_full_name() or instructor.email],
            ['Email:', instructor.email],
            ['Período:', f"{period_months} meses"],
        ]
        info_table = Table(info_data, colWidths=[2*inch, 4*inch])
        info_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f3f4f6')),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ]))
        story.append(info_table)
        story.append(Spacer(1, 0.3*inch))
        
        # Resumen general
        summary_data = [
            ['Total de Clases:', str(total_classes)],
            ['Total de Reservas:', str(total_reservations)],
            ['Total de Asistencias:', str(total_attendance)],
            ['Clases Canceladas:', str(cancelled_classes)],
            ['Ocupación Promedio:', f"{avg_occupancy:.1f}%"],
        ]
        summary_table = Table(summary_data, colWidths=[2.5*inch, 1.5*inch])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f3f4f6')),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ]))
        story.append(summary_table)
        story.append(Spacer(1, 0.3*inch))
        
        # Por tipo de clase
        if class_type_stats:
            story.append(Paragraph("Estadísticas por Tipo de Clase", styles['Heading2']))
            story.append(Spacer(1, 0.1*inch))
            
            type_data = [['Tipo', 'Cantidad', 'Reservas']]
            for stat in class_type_stats:
                type_data.append([
                    dict(Class.CLASS_TYPE_CHOICES).get(stat['class_type'], stat['class_type']),
                    str(stat['count']),
                    str(stat['total_reservations'])
                ])
            
            type_table = Table(type_data, colWidths=[2.5*inch, 1.5*inch, 2*inch])
            type_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 11),
                ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')]),
            ]))
            story.append(type_table)
            story.append(Spacer(1, 0.3*inch))
        
        # Lista de clases
        story.append(Paragraph("Clases del Período", styles['Heading2']))
        story.append(Spacer(1, 0.1*inch))
        
        classes_data = [['Fecha', 'Nombre', 'Tipo', 'Reservas', 'Capacidad']]
        for cls in classes[:50]:  # Limitar a 50 para no sobrecargar
            classes_data.append([
                cls.date.strftime('%d/%m/%Y %H:%M'),
                cls.name[:30],
                cls.get_class_type_display(),
                str(cls.reservation_count),
                str(cls.max_students)
            ])
        
        classes_table = Table(classes_data, colWidths=[1.2*inch, 2*inch, 1*inch, 0.8*inch, 0.8*inch])
        classes_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')]),
        ]))
        story.append(classes_table)
        
        if classes.count() > 50:
            story.append(Spacer(1, 0.1*inch))
            story.append(Paragraph(f"... y {classes.count() - 50} clases más", styles['Normal']))
        
        # Pie de página
        story.append(Spacer(1, 0.3*inch))
        footer = Paragraph(
            f"Generado el: {timezone.now().strftime('%d/%m/%Y %H:%M')}",
            styles['Normal']
        )
        story.append(footer)
        
        doc.build(story)
        buffer.seek(0)
        
        response = HttpResponse(buffer, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="estadisticas_instructor_{instructor.id}_{datetime.now().strftime("%Y%m%d")}.pdf"'
        return response
    
    @staticmethod
    def _generate_occupancy_pdf(start_date, end_date, occupancy_data, total_classes,
                                total_capacity, total_reservations, avg_occupancy, type_stats):
        """Genera PDF de análisis de ocupación."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        story = []
        styles = getSampleStyleSheet()
        
        # Título
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#1e40af'),
            spaceAfter=30,
            alignment=1
        )
        story.append(Paragraph("Análisis de Ocupación", title_style))
        story.append(Spacer(1, 0.2*inch))
        
        # Información del período
        info_data = [
            ['Período:', f"{start_date.strftime('%d/%m/%Y')} - {end_date.strftime('%d/%m/%Y')}"],
        ]
        info_table = Table(info_data, colWidths=[2*inch, 4*inch])
        info_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f3f4f6')),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
        ]))
        story.append(info_table)
        story.append(Spacer(1, 0.3*inch))
        
        # Resumen general
        summary_data = [
            ['Total de Clases:', str(total_classes)],
            ['Capacidad Total:', str(total_capacity)],
            ['Total de Reservas:', str(total_reservations)],
            ['Ocupación Promedio:', f"{avg_occupancy:.1f}%"],
            ['Tasa de Ocupación:', f"{(total_reservations/total_capacity*100):.1f}%" if total_capacity > 0 else "0%"],
        ]
        summary_table = Table(summary_data, colWidths=[2.5*inch, 1.5*inch])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f3f4f6')),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ]))
        story.append(summary_table)
        story.append(Spacer(1, 0.3*inch))
        
        # Por tipo de clase
        if type_stats:
            story.append(Paragraph("Ocupación por Tipo de Clase", styles['Heading2']))
            story.append(Spacer(1, 0.1*inch))
            
            type_data = [['Tipo', 'Clases', 'Reservas', 'Capacidad', 'Ocupación %']]
            for cls_type, stats in type_stats.items():
                occupancy_pct = (stats['reservations'] / stats['capacity'] * 100) if stats['capacity'] > 0 else 0
                type_data.append([
                    cls_type,
                    str(stats['total']),
                    str(stats['reservations']),
                    str(stats['capacity']),
                    f"{occupancy_pct:.1f}%"
                ])
            
            type_table = Table(type_data, colWidths=[1.5*inch, 1*inch, 1*inch, 1*inch, 1*inch])
            type_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')]),
            ]))
            story.append(type_table)
            story.append(Spacer(1, 0.3*inch))
        
        # Detalle de clases
        story.append(Paragraph("Detalle de Clases", styles['Heading2']))
        story.append(Spacer(1, 0.1*inch))
        
        detail_data = [['Fecha', 'Clase', 'Instructor', 'Reservas', 'Capacidad', 'Ocupación %']]
        for item in occupancy_data[:50]:  # Limitar a 50
            cls = item['class']
            detail_data.append([
                cls.date.strftime('%d/%m/%Y %H:%M'),
                cls.name[:25],
                cls.instructor.get_full_name()[:20] if cls.instructor else 'N/A',
                str(cls.reservation_count),
                str(cls.max_students),
                f"{item['occupancy_rate']:.1f}%"
            ])
        
        detail_table = Table(detail_data, colWidths=[1*inch, 1.5*inch, 1*inch, 0.7*inch, 0.7*inch, 0.8*inch])
        detail_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')]),
        ]))
        story.append(detail_table)
        
        if len(occupancy_data) > 50:
            story.append(Spacer(1, 0.1*inch))
            story.append(Paragraph(f"... y {len(occupancy_data) - 50} clases más", styles['Normal']))
        
        # Pie de página
        story.append(Spacer(1, 0.3*inch))
        footer = Paragraph(
            f"Generado el: {timezone.now().strftime('%d/%m/%Y %H:%M')}",
            styles['Normal']
        )
        story.append(footer)
        
        doc.build(story)
        buffer.seek(0)
        
        response = HttpResponse(buffer, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="analisis_ocupacion_{start_date.strftime("%Y%m%d")}_{end_date.strftime("%Y%m%d")}.pdf"'
        return response
    
    @staticmethod
    def _generate_waitlist_pdf(waitlist_entries, total_entries, unique_users, unique_classes, class_stats):
        """Genera PDF de reporte de lista de espera."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        story = []
        styles = getSampleStyleSheet()
        
        # Título
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#1e40af'),
            spaceAfter=30,
            alignment=1
        )
        story.append(Paragraph("Reporte de Lista de Espera", title_style))
        story.append(Spacer(1, 0.2*inch))
        
        # Resumen
        summary_data = [
            ['Total de Entradas:', str(total_entries)],
            ['Usuarios Únicos:', str(unique_users)],
            ['Clases con Lista de Espera:', str(unique_classes)],
        ]
        summary_table = Table(summary_data, colWidths=[2.5*inch, 1.5*inch])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f3f4f6')),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ]))
        story.append(summary_table)
        story.append(Spacer(1, 0.3*inch))
        
        # Por clase
        story.append(Paragraph("Lista de Espera por Clase", styles['Heading2']))
        story.append(Spacer(1, 0.1*inch))
        
        for cls_id, stats in list(class_stats.items())[:20]:  # Limitar a 20 clases
            cls = stats['class']
            story.append(Paragraph(f"{cls.name} - {cls.date.strftime('%d/%m/%Y %H:%M')}", styles['Heading3']))
            story.append(Spacer(1, 0.05*inch))
            
            entries_data = [['Posición', 'Estudiante', 'Email', 'Fecha Registro']]
            for entry in stats['entries']:
                entries_data.append([
                    str(entry.position),
                    entry.user.get_full_name() or entry.user.email,
                    entry.user.email,
                    entry.created_at.strftime('%d/%m/%Y %H:%M') if entry.created_at else 'N/A'
                ])
            
            entries_table = Table(entries_data, colWidths=[0.8*inch, 2*inch, 2.2*inch, 1.5*inch])
            entries_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 9),
                ('FONTSIZE', (0, 1), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')]),
            ]))
            story.append(entries_table)
            story.append(Spacer(1, 0.2*inch))
        
        if len(class_stats) > 20:
            story.append(Paragraph(f"... y {len(class_stats) - 20} clases más", styles['Normal']))
        
        # Pie de página
        story.append(Spacer(1, 0.3*inch))
        footer = Paragraph(
            f"Generado el: {timezone.now().strftime('%d/%m/%Y %H:%M')}",
            styles['Normal']
        )
        story.append(footer)
        
        doc.build(story)
        buffer.seek(0)
        
        response = HttpResponse(buffer, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="reporte_lista_espera_{datetime.now().strftime("%Y%m%d")}.pdf"'
        return response
    
    @staticmethod
    def _generate_cancellation_pdf(analysis):
        """Genera PDF de estadísticas de cancelaciones."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        story = []
        styles = getSampleStyleSheet()
        
        # Título
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#1e40af'),
            spaceAfter=30,
            alignment=1
        )
        story.append(Paragraph("Análisis de Cancelaciones", title_style))
        story.append(Spacer(1, 0.2*inch))
        
        # Resumen
        summary = analysis['summary']
        summary_data = [
            ['Total de Clases:', str(summary['total_classes'])],
            ['Total Canceladas:', str(summary['total_cancelled'])],
            ['Tasa de Cancelación:', f"{summary['cancellation_rate']:.1f}%"],
            ['Período:', f"{summary['period_start']} - {summary['period_end']}"],
        ]
        summary_table = Table(summary_data, colWidths=[2.5*inch, 1.5*inch])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f3f4f6')),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ]))
        story.append(summary_table)
        story.append(Spacer(1, 0.3*inch))
        
        # Por instructor
        if analysis['by_instructor']:
            story.append(Paragraph("Cancelaciones por Instructor", styles['Heading2']))
            story.append(Spacer(1, 0.1*inch))
            
            instructor_data = [['Instructor', 'Total Clases', 'Canceladas', 'Tasa %']]
            for item in analysis['by_instructor'][:10]:
                instructor_data.append([
                    item['instructor_name'],
                    str(item['total_classes']),
                    str(item['cancelled']),
                    f"{item['cancellation_rate']:.1f}%"
                ])
            
            instructor_table = Table(instructor_data, colWidths=[2.5*inch, 1.2*inch, 1.2*inch, 1*inch])
            instructor_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')]),
            ]))
            story.append(instructor_table)
            story.append(Spacer(1, 0.3*inch))
        
        # Por tipo
        if analysis['by_type']:
            story.append(Paragraph("Cancelaciones por Tipo de Clase", styles['Heading2']))
            story.append(Spacer(1, 0.1*inch))
            
            type_data = [['Tipo', 'Total', 'Canceladas', 'Tasa %']]
            for item in analysis['by_type']:
                type_data.append([
                    item['class_type'],
                    str(item['total']),
                    str(item['cancelled']),
                    f"{item['cancellation_rate']:.1f}%"
                ])
            
            type_table = Table(type_data, colWidths=[2.5*inch, 1.2*inch, 1.2*inch, 1*inch])
            type_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')]),
            ]))
            story.append(type_table)
            story.append(Spacer(1, 0.3*inch))
        
        # Razones principales
        if analysis['top_reasons']:
            story.append(Paragraph("Razones Principales de Cancelación", styles['Heading2']))
            story.append(Spacer(1, 0.1*inch))
            
            reasons_data = [['Razón', 'Cantidad']]
            for item in analysis['top_reasons'][:10]:
                reasons_data.append([
                    item['reason'] or 'Sin razón especificada',
                    str(item['count'])
                ])
            
            reasons_table = Table(reasons_data, colWidths=[4*inch, 1*inch])
            reasons_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')]),
            ]))
            story.append(reasons_table)
            story.append(Spacer(1, 0.3*inch))
        
        # Cancelaciones recientes
        if analysis['recent_cancellations']:
            story.append(Paragraph("Cancelaciones Recientes", styles['Heading2']))
            story.append(Spacer(1, 0.1*inch))
            
            recent_data = [['Fecha', 'Clase', 'Instructor', 'Razón']]
            for item in analysis['recent_cancellations'][:15]:
                recent_data.append([
                    item['cancelled_at'][:10] if item.get('cancelled_at') else 'N/A',
                    item['class_name'][:30],
                    item['instructor_name'][:20] if item.get('instructor_name') else 'N/A',
                    (item['reason'] or 'Sin razón')[:40]
                ])
            
            recent_table = Table(recent_data, colWidths=[1*inch, 2*inch, 1.5*inch, 1.5*inch])
            recent_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 9),
                ('FONTSIZE', (0, 1), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')]),
            ]))
            story.append(recent_table)
        
        # Pie de página
        story.append(Spacer(1, 0.3*inch))
        footer = Paragraph(
            f"Generado el: {timezone.now().strftime('%d/%m/%Y %H:%M')}",
            styles['Normal']
        )
        story.append(footer)
        
        doc.build(story)
        buffer.seek(0)
        
        response = HttpResponse(buffer, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="analisis_cancelaciones_{datetime.now().strftime("%Y%m%d")}.pdf"'
        return response
    
    # ========== Métodos privados para generar Excel ==========
    
    @staticmethod
    def _generate_attendance_excel(class_obj, reservations, attendance_dict):
        """Genera Excel de reporte de asistencia."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Asistencia"
        
        # Estilos
        header_fill = PatternFill(start_color="1e40af", end_color="1e40af", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        title_font = Font(bold=True, size=14, color="1e40af")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        center_align = Alignment(horizontal='center', vertical='center')
        
        # Título
        ws.merge_cells('A1:D1')
        ws['A1'] = "Reporte de Asistencia"
        ws['A1'].font = title_font
        ws['A1'].alignment = center_align
        ws.row_dimensions[1].height = 25
        
        # Información de la clase
        row = 3
        ws[f'A{row}'] = "Clase:"
        ws[f'B{row}'] = class_obj.name
        ws[f'A{row}'].font = Font(bold=True)
        row += 1
        ws[f'A{row}'] = "Fecha:"
        ws[f'B{row}'] = class_obj.date.strftime('%d/%m/%Y %H:%M')
        ws[f'A{row}'].font = Font(bold=True)
        row += 1
        ws[f'A{row}'] = "Instructor:"
        ws[f'B{row}'] = class_obj.instructor.get_full_name() if class_obj.instructor else 'N/A'
        ws[f'A{row}'].font = Font(bold=True)
        row += 1
        ws[f'A{row}'] = "Capacidad:"
        ws[f'B{row}'] = f"{class_obj.reservation_count}/{class_obj.max_students}"
        ws[f'A{row}'].font = Font(bold=True)
        row += 2
        
        # Encabezados
        headers = ['Estudiante', 'Email', 'Asistió', 'Fecha Registro']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = border
        
        # Datos
        row += 1
        for reservation in reservations:
            attended = attendance_dict.get(reservation.user_id, None)
            attended_text = 'Sí' if attended else 'No' if attended is False else 'No registrado'
            
            ws.cell(row=row, column=1, value=reservation.user.get_full_name() or reservation.user.email)
            ws.cell(row=row, column=2, value=reservation.user.email)
            ws.cell(row=row, column=3, value=attended_text)
            ws.cell(row=row, column=4, value=reservation.created_at.strftime('%d/%m/%Y %H:%M') if reservation.created_at else 'N/A')
            
            for col in range(1, 5):
                ws.cell(row=row, column=col).border = border
            
            row += 1
        
        # Resumen
        row += 1
        total = len(reservations)
        attended_count = sum(1 for v in attendance_dict.values() if v)
        no_show_count = sum(1 for v in attendance_dict.values() if v is False)
        not_recorded = total - len(attendance_dict)
        
        ws.cell(row=row, column=1, value="Total Reservas:").font = Font(bold=True)
        ws.cell(row=row, column=2, value=total)
        row += 1
        ws.cell(row=row, column=1, value="Asistieron:").font = Font(bold=True)
        ws.cell(row=row, column=2, value=attended_count)
        row += 1
        ws.cell(row=row, column=1, value="No Asistieron:").font = Font(bold=True)
        ws.cell(row=row, column=2, value=no_show_count)
        row += 1
        ws.cell(row=row, column=1, value="No Registrado:").font = Font(bold=True)
        ws.cell(row=row, column=2, value=not_recorded)
        
        # Ajustar ancho de columnas
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 18
        
        # Guardar
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="reporte_asistencia_{class_obj.id}_{datetime.now().strftime("%Y%m%d")}.xlsx"'
        return response
    
    @staticmethod
    def _generate_instructor_stats_excel(instructor, period_months, total_classes, total_reservations,
                                          total_attendance, cancelled_classes, avg_occupancy,
                                          class_type_stats, classes):
        """Genera Excel de estadísticas de instructor."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Estadísticas"
        
        # Estilos
        header_fill = PatternFill(start_color="1e40af", end_color="1e40af", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        title_font = Font(bold=True, size=14, color="1e40af")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        center_align = Alignment(horizontal='center', vertical='center')
        
        # Título
        ws.merge_cells('A1:E1')
        ws['A1'] = "Estadísticas de Instructor"
        ws['A1'].font = title_font
        ws['A1'].alignment = center_align
        ws.row_dimensions[1].height = 25
        
        # Información
        row = 3
        ws[f'A{row}'] = "Instructor:"
        ws[f'B{row}'] = instructor.get_full_name() or instructor.email
        ws[f'A{row}'].font = Font(bold=True)
        row += 1
        ws[f'A{row}'] = "Email:"
        ws[f'B{row}'] = instructor.email
        ws[f'A{row}'].font = Font(bold=True)
        row += 1
        ws[f'A{row}'] = "Período:"
        ws[f'B{row}'] = f"{period_months} meses"
        ws[f'A{row}'].font = Font(bold=True)
        row += 2
        
        # Resumen
        summary_data = [
            ['Total de Clases', total_classes],
            ['Total de Reservas', total_reservations],
            ['Total de Asistencias', total_attendance],
            ['Clases Canceladas', cancelled_classes],
            ['Ocupación Promedio', f"{avg_occupancy:.1f}%"],
        ]
        for label, value in summary_data:
            ws.cell(row=row, column=1, value=label).font = Font(bold=True)
            ws.cell(row=row, column=2, value=value)
            row += 1
        
        row += 1
        
        # Por tipo de clase
        if class_type_stats:
            ws.cell(row=row, column=1, value="Estadísticas por Tipo de Clase").font = Font(bold=True, size=12)
            row += 1
            headers = ['Tipo', 'Cantidad', 'Reservas']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col)
                cell.value = header
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align
                cell.border = border
            row += 1
            
            for stat in class_type_stats:
                ws.cell(row=row, column=1, value=dict(Class.CLASS_TYPE_CHOICES).get(stat['class_type'], stat['class_type']))
                ws.cell(row=row, column=2, value=stat['count'])
                ws.cell(row=row, column=3, value=stat['total_reservations'])
                for col in range(1, 4):
                    ws.cell(row=row, column=col).border = border
                row += 1
            
            row += 1
        
        # Lista de clases
        ws.cell(row=row, column=1, value="Clases del Período").font = Font(bold=True, size=12)
        row += 1
        headers = ['Fecha', 'Nombre', 'Tipo', 'Reservas', 'Capacidad']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = border
        row += 1
        
        for cls in classes:
            ws.cell(row=row, column=1, value=cls.date.strftime('%d/%m/%Y %H:%M'))
            ws.cell(row=row, column=2, value=cls.name)
            ws.cell(row=row, column=3, value=cls.get_class_type_display())
            ws.cell(row=row, column=4, value=cls.reservation_count)
            ws.cell(row=row, column=5, value=cls.max_students)
            for col in range(1, 6):
                ws.cell(row=row, column=col).border = border
            row += 1
        
        # Ajustar ancho de columnas
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 12
        
        # Guardar
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="estadisticas_instructor_{instructor.id}_{datetime.now().strftime("%Y%m%d")}.xlsx"'
        return response
    
    @staticmethod
    def _generate_occupancy_excel(start_date, end_date, occupancy_data, total_classes,
                                 total_capacity, total_reservations, avg_occupancy, type_stats):
        """Genera Excel de análisis de ocupación."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Ocupación"
        
        # Estilos
        header_fill = PatternFill(start_color="1e40af", end_color="1e40af", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        title_font = Font(bold=True, size=14, color="1e40af")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        center_align = Alignment(horizontal='center', vertical='center')
        
        # Título
        ws.merge_cells('A1:F1')
        ws['A1'] = "Análisis de Ocupación"
        ws['A1'].font = title_font
        ws['A1'].alignment = center_align
        ws.row_dimensions[1].height = 25
        
        # Información
        row = 3
        ws[f'A{row}'] = "Período:"
        ws[f'B{row}'] = f"{start_date.strftime('%d/%m/%Y')} - {end_date.strftime('%d/%m/%Y')}"
        ws[f'A{row}'].font = Font(bold=True)
        row += 2
        
        # Resumen
        summary_data = [
            ['Total de Clases', total_classes],
            ['Capacidad Total', total_capacity],
            ['Total de Reservas', total_reservations],
            ['Ocupación Promedio', f"{avg_occupancy:.1f}%"],
            ['Tasa de Ocupación', f"{(total_reservations/total_capacity*100):.1f}%" if total_capacity > 0 else "0%"],
        ]
        for label, value in summary_data:
            ws.cell(row=row, column=1, value=label).font = Font(bold=True)
            ws.cell(row=row, column=2, value=value)
            row += 1
        
        row += 1
        
        # Por tipo
        if type_stats:
            ws.cell(row=row, column=1, value="Ocupación por Tipo de Clase").font = Font(bold=True, size=12)
            row += 1
            headers = ['Tipo', 'Clases', 'Reservas', 'Capacidad', 'Ocupación %']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col)
                cell.value = header
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align
                cell.border = border
            row += 1
            
            for cls_type, stats in type_stats.items():
                occupancy_pct = (stats['reservations'] / stats['capacity'] * 100) if stats['capacity'] > 0 else 0
                ws.cell(row=row, column=1, value=cls_type)
                ws.cell(row=row, column=2, value=stats['total'])
                ws.cell(row=row, column=3, value=stats['reservations'])
                ws.cell(row=row, column=4, value=stats['capacity'])
                ws.cell(row=row, column=5, value=f"{occupancy_pct:.1f}%")
                for col in range(1, 6):
                    ws.cell(row=row, column=col).border = border
                row += 1
            
            row += 1
        
        # Detalle
        ws.cell(row=row, column=1, value="Detalle de Clases").font = Font(bold=True, size=12)
        row += 1
        headers = ['Fecha', 'Clase', 'Instructor', 'Reservas', 'Capacidad', 'Ocupación %']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = border
        row += 1
        
        for item in occupancy_data:
            cls = item['class']
            ws.cell(row=row, column=1, value=cls.date.strftime('%d/%m/%Y %H:%M'))
            ws.cell(row=row, column=2, value=cls.name)
            ws.cell(row=row, column=3, value=cls.instructor.get_full_name() if cls.instructor else 'N/A')
            ws.cell(row=row, column=4, value=cls.reservation_count)
            ws.cell(row=row, column=5, value=cls.max_students)
            ws.cell(row=row, column=6, value=f"{item['occupancy_rate']:.1f}%")
            for col in range(1, 7):
                ws.cell(row=row, column=col).border = border
            row += 1
        
        # Ajustar ancho de columnas
        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 20
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 12
        ws.column_dimensions['F'].width = 15
        
        # Guardar
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="analisis_ocupacion_{start_date.strftime("%Y%m%d")}_{end_date.strftime("%Y%m%d")}.xlsx"'
        return response
    
    @staticmethod
    def _generate_waitlist_excel(waitlist_entries, total_entries, unique_users, unique_classes, class_stats):
        """Genera Excel de reporte de lista de espera."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Lista de Espera"
        
        # Estilos
        header_fill = PatternFill(start_color="1e40af", end_color="1e40af", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        title_font = Font(bold=True, size=14, color="1e40af")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        center_align = Alignment(horizontal='center', vertical='center')
        
        # Título
        ws.merge_cells('A1:D1')
        ws['A1'] = "Reporte de Lista de Espera"
        ws['A1'].font = title_font
        ws['A1'].alignment = center_align
        ws.row_dimensions[1].height = 25
        
        # Resumen
        row = 3
        summary_data = [
            ['Total de Entradas', total_entries],
            ['Usuarios Únicos', unique_users],
            ['Clases con Lista de Espera', unique_classes],
        ]
        for label, value in summary_data:
            ws.cell(row=row, column=1, value=label).font = Font(bold=True)
            ws.cell(row=row, column=2, value=value)
            row += 1
        
        row += 1
        
        # Por clase
        for cls_id, stats in class_stats.items():
            cls = stats['class']
            ws.cell(row=row, column=1, value=f"{cls.name} - {cls.date.strftime('%d/%m/%Y %H:%M')}").font = Font(bold=True, size=12)
            row += 1
            
            headers = ['Posición', 'Estudiante', 'Email', 'Fecha Registro']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col)
                cell.value = header
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align
                cell.border = border
            row += 1
            
            for entry in stats['entries']:
                ws.cell(row=row, column=1, value=entry.position)
                ws.cell(row=row, column=2, value=entry.user.get_full_name() or entry.user.email)
                ws.cell(row=row, column=3, value=entry.user.email)
                ws.cell(row=row, column=4, value=entry.created_at.strftime('%d/%m/%Y %H:%M') if entry.created_at else 'N/A')
                for col in range(1, 5):
                    ws.cell(row=row, column=col).border = border
                row += 1
            
            row += 1
        
        # Ajustar ancho de columnas
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 30
        ws.column_dimensions['D'].width = 18
        
        # Guardar
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="reporte_lista_espera_{datetime.now().strftime("%Y%m%d")}.xlsx"'
        return response
    
    @staticmethod
    def _generate_cancellation_excel(analysis):
        """Genera Excel de estadísticas de cancelaciones."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Cancelaciones"
        
        # Estilos
        header_fill = PatternFill(start_color="1e40af", end_color="1e40af", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        title_font = Font(bold=True, size=14, color="1e40af")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        center_align = Alignment(horizontal='center', vertical='center')
        
        # Título
        ws.merge_cells('A1:D1')
        ws['A1'] = "Análisis de Cancelaciones"
        ws['A1'].font = title_font
        ws['A1'].alignment = center_align
        ws.row_dimensions[1].height = 25
        
        # Resumen
        summary = analysis['summary']
        row = 3
        summary_data = [
            ['Total de Clases', summary['total_classes']],
            ['Total Canceladas', summary['total_cancelled']],
            ['Tasa de Cancelación', f"{summary['cancellation_rate']:.1f}%"],
            ['Período', f"{summary['period_start']} - {summary['period_end']}"],
        ]
        for label, value in summary_data:
            ws.cell(row=row, column=1, value=label).font = Font(bold=True)
            ws.cell(row=row, column=2, value=value)
            row += 1
        
        row += 1
        
        # Por instructor
        if analysis['by_instructor']:
            ws.cell(row=row, column=1, value="Cancelaciones por Instructor").font = Font(bold=True, size=12)
            row += 1
            headers = ['Instructor', 'Total Clases', 'Canceladas', 'Tasa %']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col)
                cell.value = header
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align
                cell.border = border
            row += 1
            
            for item in analysis['by_instructor']:
                ws.cell(row=row, column=1, value=item['instructor_name'])
                ws.cell(row=row, column=2, value=item['total_classes'])
                ws.cell(row=row, column=3, value=item['cancelled'])
                ws.cell(row=row, column=4, value=f"{item['cancellation_rate']:.1f}%")
                for col in range(1, 5):
                    ws.cell(row=row, column=col).border = border
                row += 1
            
            row += 1
        
        # Por tipo
        if analysis['by_type']:
            ws.cell(row=row, column=1, value="Cancelaciones por Tipo de Clase").font = Font(bold=True, size=12)
            row += 1
            headers = ['Tipo', 'Total', 'Canceladas', 'Tasa %']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col)
                cell.value = header
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align
                cell.border = border
            row += 1
            
            for item in analysis['by_type']:
                ws.cell(row=row, column=1, value=item['class_type'])
                ws.cell(row=row, column=2, value=item['total'])
                ws.cell(row=row, column=3, value=item['cancelled'])
                ws.cell(row=row, column=4, value=f"{item['cancellation_rate']:.1f}%")
                for col in range(1, 5):
                    ws.cell(row=row, column=col).border = border
                row += 1
            
            row += 1
        
        # Razones principales
        if analysis['top_reasons']:
            ws.cell(row=row, column=1, value="Razones Principales de Cancelación").font = Font(bold=True, size=12)
            row += 1
            headers = ['Razón', 'Cantidad']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col)
                cell.value = header
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align
                cell.border = border
            row += 1
            
            for item in analysis['top_reasons']:
                ws.cell(row=row, column=1, value=item['reason'] or 'Sin razón especificada')
                ws.cell(row=row, column=2, value=item['count'])
                for col in range(1, 3):
                    ws.cell(row=row, column=col).border = border
                row += 1
            
            row += 1
        
        # Cancelaciones recientes
        if analysis['recent_cancellations']:
            ws.cell(row=row, column=1, value="Cancelaciones Recientes").font = Font(bold=True, size=12)
            row += 1
            headers = ['Fecha', 'Clase', 'Instructor', 'Razón']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col)
                cell.value = header
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align
                cell.border = border
            row += 1
            
            for item in analysis['recent_cancellations']:
                ws.cell(row=row, column=1, value=item['cancelled_at'][:10] if item.get('cancelled_at') else 'N/A')
                ws.cell(row=row, column=2, value=item['class_name'])
                ws.cell(row=row, column=3, value=item.get('instructor_name', 'N/A'))
                ws.cell(row=row, column=4, value=item['reason'] or 'Sin razón')
                for col in range(1, 5):
                    ws.cell(row=row, column=col).border = border
                row += 1
        
        # Ajustar ancho de columnas
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 20
        ws.column_dimensions['D'].width = 40
        
        # Guardar
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="analisis_cancelaciones_{datetime.now().strftime("%Y%m%d")}.xlsx"'
        return response

