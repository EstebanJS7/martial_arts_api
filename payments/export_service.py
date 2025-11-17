"""
Servicio para generar reportes exportables de pagos en PDF y Excel.
"""
import io
from datetime import datetime, timedelta
from django.http import HttpResponse
from django.utils import timezone
from django.db.models import Count, Sum, Q, F
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from .models import Payment, PaymentTransaction
from django.contrib.auth import get_user_model
from decimal import Decimal

User = get_user_model()


class PaymentExportService:
    """
    Servicio para exportar reportes de pagos en PDF y Excel.
    """
    
    @staticmethod
    def export_payment_status_report(start_date=None, end_date=None, format='pdf'):
        """
        Exporta reporte del estado de pagos de usuarios.
        Muestra quiénes están al día, con deuda, etc.
        
        Args:
            start_date: Fecha de inicio (opcional)
            end_date: Fecha de fin (opcional)
            format: 'pdf' o 'excel'
        """
        today = timezone.now().date()
        
        # Filtrar pagos
        payments_qs = Payment.objects.select_related('user').all()
        
        if start_date:
            payments_qs = payments_qs.filter(due_date__gte=start_date)
        if end_date:
            payments_qs = payments_qs.filter(due_date__lte=end_date)
        
        # Agrupar por usuario y calcular estado
        users_data = {}
        for payment in payments_qs:
            user_id = payment.user_id
            if user_id not in users_data:
                users_data[user_id] = {
                    'user': payment.user,
                    'total_amount': Decimal('0'),
                    'total_paid': Decimal('0'),
                    'total_pending': Decimal('0'),
                    'overdue_amount': Decimal('0'),
                    'overdue_count': 0,
                    'pending_count': 0,
                    'paid_count': 0,
                    'payments': []
                }
            
            user_data = users_data[user_id]
            user_data['total_amount'] += payment.amount
            user_data['total_paid'] += payment.amount_paid
            user_data['total_pending'] += (payment.amount - payment.amount_paid)
            
            if payment.is_fully_paid:
                user_data['paid_count'] += 1
            elif payment.due_date < today and not payment.is_fully_paid:
                user_data['overdue_count'] += 1
                user_data['overdue_amount'] += (payment.amount - payment.amount_paid)
            else:
                user_data['pending_count'] += 1
            
            user_data['payments'].append(payment)
        
        # Clasificar usuarios por estado
        users_by_status = {
            'al_dia': [],  # Sin deudas, todos los pagos al día
            'con_deuda': [],  # Tiene pagos vencidos
            'pendientes': [],  # Tiene pagos pendientes pero no vencidos
            'al_dia_parcial': [],  # Al día pero con pagos parciales
        }
        
        for user_id, data in users_data.items():
            if data['overdue_count'] > 0:
                status = 'con_deuda'
            elif data['pending_count'] > 0:
                status = 'pendientes'
            elif data['total_paid'] < data['total_amount']:
                status = 'al_dia_parcial'
            else:
                status = 'al_dia'
            
            users_by_status[status].append(data)
        
        # Calcular resumen
        total_users = len(users_data)
        total_amount = sum(d['total_amount'] for d in users_data.values())
        total_paid = sum(d['total_paid'] for d in users_data.values())
        total_pending = sum(d['total_pending'] for d in users_data.values())
        total_overdue = sum(d['overdue_amount'] for d in users_data.values())
        
        summary = {
            'total_users': total_users,
            'al_dia': len(users_by_status['al_dia']),
            'con_deuda': len(users_by_status['con_deuda']),
            'pendientes': len(users_by_status['pendientes']),
            'al_dia_parcial': len(users_by_status['al_dia_parcial']),
            'total_amount': total_amount,
            'total_paid': total_paid,
            'total_pending': total_pending,
            'total_overdue': total_overdue,
            'collection_rate': (total_paid / total_amount * 100) if total_amount > 0 else 0,
        }
        
        if format == 'pdf':
            return PaymentExportService._generate_payment_status_pdf(
                users_by_status, summary, start_date, end_date
            )
        else:
            return PaymentExportService._generate_payment_status_excel(
                users_by_status, summary, start_date, end_date
            )
    
    # ========== Métodos privados para generar PDF ==========
    
    @staticmethod
    def _generate_payment_status_pdf(users_by_status, summary, start_date, end_date):
        """Genera PDF de reporte de estado de pagos."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        story = []
        styles = getSampleStyleSheet()
        
        # Título
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#1e40af'),
            spaceAfter=30,
            alignment=1
        )
        story.append(Paragraph("Reporte de Estado de Pagos", title_style))
        story.append(Spacer(1, 0.2*inch))
        
        # Información del período
        period_text = "Todos los pagos"
        if start_date or end_date:
            period_parts = []
            if start_date:
                period_parts.append(f"Desde: {start_date.strftime('%d/%m/%Y')}")
            if end_date:
                period_parts.append(f"Hasta: {end_date.strftime('%d/%m/%Y')}")
            period_text = " - ".join(period_parts)
        
        info_data = [
            ['Período:', period_text],
        ]
        info_table = Table(info_data, colWidths=[2*inch, 4*inch])
        info_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f3f4f6')),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
        ]))
        story.append(info_table)
        story.append(Spacer(1, 0.3*inch))
        
        # Resumen general
        story.append(Paragraph("Resumen General", styles['Heading2']))
        story.append(Spacer(1, 0.1*inch))
        
        summary_data = [
            ['Total de Usuarios:', str(summary['total_users'])],
            ['Usuarios al Día:', str(summary['al_dia'])],
            ['Usuarios con Deuda:', str(summary['con_deuda'])],
            ['Usuarios con Pendientes:', str(summary['pendientes'])],
            ['Usuarios al Día (Parciales):', str(summary['al_dia_parcial'])],
            ['Monto Total:', f"${summary['total_amount']:,.2f}"],
            ['Monto Pagado:', f"${summary['total_paid']:,.2f}"],
            ['Monto Pendiente:', f"${summary['total_pending']:,.2f}"],
            ['Monto Vencido:', f"${summary['total_overdue']:,.2f}"],
            ['Tasa de Recaudación:', f"{summary['collection_rate']:.1f}%"],
        ]
        summary_table = Table(summary_data, colWidths=[2.5*inch, 1.5*inch])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f3f4f6')),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ]))
        story.append(summary_table)
        story.append(Spacer(1, 0.3*inch))
        
        # Usuarios con deuda
        if users_by_status['con_deuda']:
            story.append(Paragraph("Usuarios con Deuda", styles['Heading2']))
            story.append(Spacer(1, 0.1*inch))
            
            debt_data = [['Usuario', 'Email', 'Monto Total', 'Pagado', 'Deuda', 'Pagos Vencidos']]
            for user_data in sorted(users_by_status['con_deuda'], key=lambda x: x['overdue_amount'], reverse=True)[:50]:
                user = user_data['user']
                debt_data.append([
                    user.get_full_name() or user.email,
                    user.email,
                    f"${user_data['total_amount']:,.2f}",
                    f"${user_data['total_paid']:,.2f}",
                    f"${user_data['overdue_amount']:,.2f}",
                    str(user_data['overdue_count'])
                ])
            
            debt_table = Table(debt_data, colWidths=[1.5*inch, 2*inch, 1*inch, 1*inch, 1*inch, 1*inch])
            debt_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#dc2626')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 9),
                ('FONTSIZE', (0, 1), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#fef2f2')]),
            ]))
            story.append(debt_table)
            story.append(Spacer(1, 0.3*inch))
        
        # Usuarios pendientes
        if users_by_status['pendientes']:
            story.append(Paragraph("Usuarios con Pagos Pendientes", styles['Heading2']))
            story.append(Spacer(1, 0.1*inch))
            
            pending_data = [['Usuario', 'Email', 'Monto Total', 'Pagado', 'Pendiente', 'Pagos Pendientes']]
            for user_data in sorted(users_by_status['pendientes'], key=lambda x: x['total_pending'], reverse=True)[:30]:
                user = user_data['user']
                pending_data.append([
                    user.get_full_name() or user.email,
                    user.email,
                    f"${user_data['total_amount']:,.2f}",
                    f"${user_data['total_paid']:,.2f}",
                    f"${user_data['total_pending']:,.2f}",
                    str(user_data['pending_count'])
                ])
            
            pending_table = Table(pending_data, colWidths=[1.5*inch, 2*inch, 1*inch, 1*inch, 1*inch, 1*inch])
            pending_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f59e0b')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 9),
                ('FONTSIZE', (0, 1), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#fffbeb')]),
            ]))
            story.append(pending_table)
            story.append(Spacer(1, 0.3*inch))
        
        # Usuarios al día
        if users_by_status['al_dia']:
            story.append(Paragraph("Usuarios al Día", styles['Heading2']))
            story.append(Spacer(1, 0.1*inch))
            
            current_data = [['Usuario', 'Email', 'Monto Total', 'Pagado', 'Pagos Completados']]
            for user_data in sorted(users_by_status['al_dia'], key=lambda x: x['total_paid'], reverse=True)[:30]:
                user = user_data['user']
                current_data.append([
                    user.get_full_name() or user.email,
                    user.email,
                    f"${user_data['total_amount']:,.2f}",
                    f"${user_data['total_paid']:,.2f}",
                    str(user_data['paid_count'])
                ])
            
            current_table = Table(current_data, colWidths=[1.5*inch, 2*inch, 1*inch, 1*inch, 1.5*inch])
            current_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#10b981')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 9),
                ('FONTSIZE', (0, 1), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f0fdf4')]),
            ]))
            story.append(current_table)
        
        # Pie de página
        story.append(Spacer(1, 0.3*inch))
        footer = Paragraph(
            f"Generado el: {timezone.now().strftime('%d/%m/%Y %H:%M')}",
            styles['Normal']
        )
        story.append(footer)
        
        doc.build(story)
        buffer.seek(0)
        
        filename = f"reporte_estado_pagos_{datetime.now().strftime('%Y%m%d')}.pdf"
        if start_date and end_date:
            filename = f"reporte_estado_pagos_{start_date.strftime('%Y%m%d')}_{end_date.strftime('%Y%m%d')}.pdf"
        
        response = HttpResponse(buffer, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
    
    # ========== Métodos privados para generar Excel ==========
    
    @staticmethod
    def _generate_payment_status_excel(users_by_status, summary, start_date, end_date):
        """Genera Excel de reporte de estado de pagos."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Estado de Pagos"
        
        # Estilos
        header_fill = PatternFill(start_color="1e40af", end_color="1e40af", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        title_font = Font(bold=True, size=14, color="1e40af")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        center_align = Alignment(horizontal='center', vertical='center')
        
        # Título
        ws.merge_cells('A1:F1')
        ws['A1'] = "Reporte de Estado de Pagos"
        ws['A1'].font = title_font
        ws['A1'].alignment = center_align
        ws.row_dimensions[1].height = 25
        
        # Información del período
        row = 3
        period_text = "Todos los pagos"
        if start_date or end_date:
            period_parts = []
            if start_date:
                period_parts.append(f"Desde: {start_date.strftime('%d/%m/%Y')}")
            if end_date:
                period_parts.append(f"Hasta: {end_date.strftime('%d/%m/%Y')}")
            period_text = " - ".join(period_parts)
        
        ws.cell(row=row, column=1, value="Período:").font = Font(bold=True)
        ws.cell(row=row, column=2, value=period_text)
        row += 2
        
        # Resumen
        ws.cell(row=row, column=1, value="Resumen General").font = Font(bold=True, size=12)
        row += 1
        summary_data = [
            ['Total de Usuarios', summary['total_users']],
            ['Usuarios al Día', summary['al_dia']],
            ['Usuarios con Deuda', summary['con_deuda']],
            ['Usuarios con Pendientes', summary['pendientes']],
            ['Usuarios al Día (Parciales)', summary['al_dia_parcial']],
            ['Monto Total', f"${summary['total_amount']:,.2f}"],
            ['Monto Pagado', f"${summary['total_paid']:,.2f}"],
            ['Monto Pendiente', f"${summary['total_pending']:,.2f}"],
            ['Monto Vencido', f"${summary['total_overdue']:,.2f}"],
            ['Tasa de Recaudación', f"{summary['collection_rate']:.1f}%"],
        ]
        for label, value in summary_data:
            ws.cell(row=row, column=1, value=label).font = Font(bold=True)
            ws.cell(row=row, column=2, value=value)
            row += 1
        
        row += 1
        
        # Usuarios con deuda
        if users_by_status['con_deuda']:
            ws.cell(row=row, column=1, value="Usuarios con Deuda").font = Font(bold=True, size=12)
            row += 1
            headers = ['Usuario', 'Email', 'Monto Total', 'Pagado', 'Deuda', 'Pagos Vencidos']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col)
                cell.value = header
                cell.fill = PatternFill(start_color="dc2626", end_color="dc2626", fill_type="solid")
                cell.font = header_font
                cell.alignment = center_align
                cell.border = border
            row += 1
            
            for user_data in sorted(users_by_status['con_deuda'], key=lambda x: x['overdue_amount'], reverse=True):
                user = user_data['user']
                ws.cell(row=row, column=1, value=user.get_full_name() or user.email)
                ws.cell(row=row, column=2, value=user.email)
                ws.cell(row=row, column=3, value=f"${user_data['total_amount']:,.2f}")
                ws.cell(row=row, column=4, value=f"${user_data['total_paid']:,.2f}")
                ws.cell(row=row, column=5, value=f"${user_data['overdue_amount']:,.2f}")
                ws.cell(row=row, column=6, value=user_data['overdue_count'])
                for col in range(1, 7):
                    ws.cell(row=row, column=col).border = border
                row += 1
            
            row += 1
        
        # Usuarios pendientes
        if users_by_status['pendientes']:
            ws.cell(row=row, column=1, value="Usuarios con Pagos Pendientes").font = Font(bold=True, size=12)
            row += 1
            headers = ['Usuario', 'Email', 'Monto Total', 'Pagado', 'Pendiente', 'Pagos Pendientes']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col)
                cell.value = header
                cell.fill = PatternFill(start_color="f59e0b", end_color="f59e0b", fill_type="solid")
                cell.font = header_font
                cell.alignment = center_align
                cell.border = border
            row += 1
            
            for user_data in sorted(users_by_status['pendientes'], key=lambda x: x['total_pending'], reverse=True):
                user = user_data['user']
                ws.cell(row=row, column=1, value=user.get_full_name() or user.email)
                ws.cell(row=row, column=2, value=user.email)
                ws.cell(row=row, column=3, value=f"${user_data['total_amount']:,.2f}")
                ws.cell(row=row, column=4, value=f"${user_data['total_paid']:,.2f}")
                ws.cell(row=row, column=5, value=f"${user_data['total_pending']:,.2f}")
                ws.cell(row=row, column=6, value=user_data['pending_count'])
                for col in range(1, 7):
                    ws.cell(row=row, column=col).border = border
                row += 1
            
            row += 1
        
        # Usuarios al día
        if users_by_status['al_dia']:
            ws.cell(row=row, column=1, value="Usuarios al Día").font = Font(bold=True, size=12)
            row += 1
            headers = ['Usuario', 'Email', 'Monto Total', 'Pagado', 'Pagos Completados']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col)
                cell.value = header
                cell.fill = PatternFill(start_color="10b981", end_color="10b981", fill_type="solid")
                cell.font = header_font
                cell.alignment = center_align
                cell.border = border
            row += 1
            
            for user_data in sorted(users_by_status['al_dia'], key=lambda x: x['total_paid'], reverse=True):
                user = user_data['user']
                ws.cell(row=row, column=1, value=user.get_full_name() or user.email)
                ws.cell(row=row, column=2, value=user.email)
                ws.cell(row=row, column=3, value=f"${user_data['total_amount']:,.2f}")
                ws.cell(row=row, column=4, value=f"${user_data['total_paid']:,.2f}")
                ws.cell(row=row, column=5, value=user_data['paid_count'])
                for col in range(1, 6):
                    ws.cell(row=row, column=col).border = border
                row += 1
        
        # Ajustar ancho de columnas
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 15
        ws.column_dimensions['F'].width = 18
        
        # Guardar
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        filename = f"reporte_estado_pagos_{datetime.now().strftime('%Y%m%d')}.xlsx"
        if start_date and end_date:
            filename = f"reporte_estado_pagos_{start_date.strftime('%Y%m%d')}_{end_date.strftime('%Y%m%d')}.xlsx"
        
        response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

